"""
configuration.py
================
Python port of the DataConf case-1 helper function files:
    configurationNewData.m, generateBuildingConf.m,
    generateCarConf.m, generateTrafficConf.m
data_conf.DataConf imports these lazily for dataType 1, plus
configurationNewDataWithCustomInitials (dataType 3, 'New Traffic with
Custom Initials') ported from matlab_src/utils/.

CERTIFIED against matlab_src/utils/ originals (uploaded 2026-07):
configurationNewData and generateBuildingConf match line for line;
generateCarConf confirms [F2b] (its `parkAlgorithm` loop variable is
never written onto the cars) and [F3] (stale `cars` array reused
across configurations); generateTrafficConf confirms [F1] (hardcoded
`:10:` grid step) and the Traffic.empty placeholder that the port
stores as None.

Deviations / fixes ([F1]..[F3] inline):

  [F1] generateTrafficConf hardcodes a 10-percentage-point grid step
       (INCmin:10:INCmax). Exposed as TRAFFIC_STEP so that
       Experiment.saveResults labels rows/columns with the same step
       (see experiment.py [D8]).

  [F2] The parking-algorithm chain was broken in MATLAB three ways:
       (a) the GUI put the DropDown *widget handle* into
           startData.parkingAlgorithm, not its value;
       (b) the loop variable `parkAlgorithm` was never written onto the
           cars, so setParkFloor compared against Car's default (1);
       (c) consequently Park2..Park4 never changed anything.
       The port parses the GUI string ("Park2" -> 2, "Park1 (No parking
       method)" -> 1), assigns it to each car, then applies
       setParkFloor - so parking selections behave as evidently
       intended. Park1 remains "no parking" (parkFloor stays 'N/A').

  [F3] MATLAB reused a stale function-local `cars` array across
       configurations (safe only because NCmin:NCmax ascends); the port
       builds a fresh list per configuration.
"""
from __future__ import annotations

import re
from typing import Any, List, Optional

from building import Building
from car import Car
from data_structures import HallCallLists, PassengerLists
from data_structures.hall_call import HallCall
from data_structures.passenger import Passenger
from traffic import Traffic

TRAFFIC_STEP = 10  # [F1] MATLAB: INCmin:10:INCmax / INTmin:10:INTmax


def configurationNewData(startData: Any, dataConf: Any) -> None:
    """Port of configurationNewData.m (in-place population of dataConf)."""
    # Building configuration
    dataConf.BUILDING = generateBuildingConf(startData)
    # Simulation configuration: `data.simulation = Simulation(startData)`
    # was commented out in the MATLAB source - kept out here too.
    # Car configuration
    dataConf.CAR = generateCarConf(startData)
    # Traffic configuration
    dataConf.TRAFFIC = generateTrafficConf(startData)


def configurationNewDataWithCustomInitials(startData: Any, dataConf: Any) -> None:
    """
    Port of matlab_src/utils/configurationNewDataWithCustomInitials.m
    (dataType 3, 'New Traffic with Custom Initials'): populate dataConf
    from an initials workbook (matlab_src/initials_file.xlsx layout).

    Workbook layout - 16 columns, 2 header rows, data from row 3 (the
    15-column comment block inside the .m is stale; the code reads):
      A nf | B incoming% | C interfloor% | D outgoing% (read, unused)
      E nc | F parkAlg | G carInitials ('true' -> apply H..J)
      H car floors | I car states | J car DFs (number or 'd1 d2 ..')
      K/L/M up-passenger floor / DF / assigned car
      N/O/P down-passenger floor / DF / assigned car

    The scalars OVERRIDE the Building & Car / Traffic tab grid: one
    building, one traffic cell, one car set. Waiting passengers are
    floor-sorted per direction with one hall call per unique floor
    (assigned-car columns empty -> carId 0, isInitialDispatch False);
    every car destination floor becomes an already-boarded passenger
    (QJT = BT = -1, load incremented) - Simulator.fillResultsTable
    excludes those from the dataType-3 statistics via QJT != -1.

    Deviations ([F4]..[F7]; the .m's debug echo of P_initials(i).floor
    and its unused filterNumeric lambda are omitted):

      [F4] readcell -> openpyxl values: MATLAB 'missing' cells arrive
           as None and are dropped by the same numeric-mask rule as the
           validIdx filters (bools excluded - MATLAB logicals are not
           isnumeric). Out-of-range car-DF rows yield [] exactly like
           the .m's try/catch/[] ladder.
      [F5] carInitials: the .m compares == "true" (case-sensitive, text
           cells only; an Excel boolean TRUE would NOT activate it in
           MATLAB). The port matches the text 'true' and, leniently,
           accepts a real boolean True from openpyxl.
      [F6] Car.DF is a set in the port [P19]: the boarded in-car
           passengers are created in ascending DF order (the .m
           followed the file's cell order, which a set cannot
           preserve; only Passenger ids / travelling order differ).
      [F7] certified against ESRA_v2 Car.m: its constructor never reads
           startData.parkAlg, so the file's Park Algorithm was DEAD in
           MATLAB (setParkFloor is not called here either). The port
           stores it on startData.parkAlg like the .m and also onto
           car.parkAlgorithm for inspection - behaviour-identical
           today because parkFloor stays 'N/A' either way.
    """
    import openpyxl                    # lazy: needed only for dataType 3

    wb = openpyxl.load_workbook(startData.fileName, data_only=True)
    ws = wb[wb.sheetnames[0]]
    # table = table(3:end,:) - data rows, padded to the 16-column layout
    rows = [(list(r) + [None] * 16)[:16]
            for r in ws.iter_rows(min_row=3, values_only=True)]
    if not rows:
        raise ValueError(f"initials file '{startData.fileName}' has no "
                         "data rows (data starts at spreadsheet row 3).")

    def numericColumn(col: int) -> List[Any]:
        # [F4] MATLAB validIdx: isnumeric(x) && ~ismissing(x)
        return [row[col] for row in rows
                if isinstance(row[col], (int, float))
                and not isinstance(row[col], bool)]

    head = rows[0]
    nf = int(head[0])
    inc, int_ = int(head[1]), int(head[2])
    # head[3] (outgoing %) is read but unused in the original as well
    nc = int(head[4])
    parkAlg = int(head[5])
    carInitials = head[6]

    dataConf.BUILDING = [Building(nf)]
    dataConf.TRAFFIC = [[Traffic(inc, int_)]]
    # (unlike generateTrafficConf, the .m does NOT push
    #  numInitialPassengers into Traffic here - preserved)

    startData.nc = nc
    startData.parkAlg = parkAlg                              # [F7]
    startData.carInitials = carInitials
    cars = [Car(startData, car_id) for car_id in range(1, nc + 1)]
    for car in cars:
        car.parkAlgorithm = parkAlg                          # [F7]

    applyInitials = (carInitials is True                     # [F5]
                     or (isinstance(carInitials, str)
                         and carInitials.strip() == "true"))
    if applyInitials:
        carFloors = numericColumn(7)
        carStates = numericColumn(8)
        if len(carFloors) < nc or len(carStates) < nc:
            raise ValueError(
                f"initials file: 'Car floors'/'Car states' columns hold "
                f"fewer than nc={nc} numeric values.")
        for i, car in enumerate(cars):
            raw = rows[i][9] if i < len(rows) else None      # [F4]
            if isinstance(raw, str):
                DF = [int(tok) for tok in raw.split()]       # sscanf '%d '
            elif isinstance(raw, (int, float)) and not isinstance(raw, bool):
                DF = [int(raw)]
            else:
                DF = []
            car.floor = float(carFloors[i])
            car.previousFloor = car.floor
            car.state = int(carStates[i])
            car.DF = set(DF)

    P = PassengerLists()
    HC = HallCallLists()
    anyAssigned = False

    for dir_ in (1, 2):                  # 1: up cols K-M, 2: down N-P
        base = 10 + (dir_ - 1) * 3
        floors = numericColumn(base)
        dfs = numericColumn(base + 1)
        assigned = numericColumn(base + 2)
        if len(dfs) != len(floors) or (assigned
                                       and len(assigned) != len(floors)):
            raise ValueError(
                f"initials file: direction-{dir_} passenger columns are "
                f"misaligned ({len(floors)} floors, {len(dfs)} DFs, "
                f"{len(assigned)} assigned cars).")

        order = sorted(range(len(floors)), key=lambda k: floors[k])
        floors = [floors[k] for k in order]      # stable, like sort()
        dfs = [dfs[k] for k in order]
        if assigned:
            assigned = [assigned[k] for k in order]
            anyAssigned = True

        for j, floor in enumerate(floors):
            passenger = Passenger(int(floor), 0, int(dfs[j]))
            if assigned:               # else carId stays the 0 default
                passenger.carId = int(assigned[j])
            P.add(passenger)

        seen = set()
        for j, floor in enumerate(floors):   # unique() first occurrences
            if floor not in seen:
                seen.add(floor)
                hc = HallCall(int(floor), 0, dir_)
                hc.carId = int(assigned[j]) if assigned else 0
                HC.add(hc)

    # 1 if cars already assigned, 0 if not (either direction suffices)
    dataConf.isInitialDispatch = anyAssigned

    for car in cars:                     # car DFs -> boarded passengers
        for df in sorted(car.DF):                            # [F6]
            passenger = Passenger(int(car.floor), -1, int(df))
            passenger.board(car.id, -1)
            car.P.transfer(passenger, "travelling")
            car.load += 1

    dataConf.CAR = [cars]
    dataConf.initialCars = cars          # same objects, as in MATLAB
    dataConf.initialP = P
    dataConf.initialHC = HC


def generateBuildingConf(startData: Any) -> List[Building]:
    """Port of generateBuildingConf.m: one Building per NFmin:NFstep:NFmax."""
    return [Building(n_floor)
            for n_floor in range(int(startData.NFmin),
                                 int(startData.NFmax) + 1,
                                 int(startData.NFstep))]


def generateCarConf(startData: Any) -> List[List[Car]]:
    """
    Port of generateCarConf.m: one car list per
    (number of cars) x (parking algorithm) combination.
    """
    parkingAlgorithms = _parseParkingAlgorithms(startData.parkingAlgorithm)
    CAR: List[List[Car]] = []
    for n_car in range(int(startData.NCmin), int(startData.NCmax) + 1):
        for parkAlgorithm in parkingAlgorithms:
            cars: List[Car] = []                     # [F3] fresh list
            for car_id in range(1, n_car + 1):
                car = Car(startData, car_id)
                car.parkAlgorithm = parkAlgorithm    # [F2b] actually applied
                _setParkFloor(car)
                cars.append(car)
            CAR.append(cars)
    return CAR


def _setParkFloor(car: Car) -> None:
    """Port of the nested setParkFloor function in generateCarConf.m."""
    if car.parkAlgorithm == 2:
        car.parkFloor = 1


def generateTrafficConf(startData: Any) -> List[List[Optional[Traffic]]]:
    """
    Port of generateTrafficConf.m: rows = incoming %, cols = interfloor %,
    both stepping by TRAFFIC_STEP. Combinations with inc + int > 100 are
    invalid and stored as None (MATLAB: Traffic.empty) - Experiment.run
    skips them and saveResults reports NaN for them.
    """
    Traffic.getSetNumInitialPassengers(startData.numInitialPassengers)

    incValues = range(int(startData.INCmin), int(startData.INCmax) + 1,
                      TRAFFIC_STEP)
    intValues = range(int(startData.INTmin), int(startData.INTmax) + 1,
                      TRAFFIC_STEP)

    TRAFFIC: List[List[Optional[Traffic]]] = []
    for inc in incValues:
        row: List[Optional[Traffic]] = []
        for int_ in intValues:
            row.append(Traffic(inc, int_) if (100 - inc - int_) >= 0 else None)
        TRAFFIC.append(row)
    return TRAFFIC


def _parseParkingAlgorithms(value: Any) -> List[int]:
    """
    [F2a] Normalizes startData.parkingAlgorithm into a list of ints.
    Accepts an int, a list/tuple of ints or strings, or GUI strings such
    as 'Park2' / 'Park1 (No parking method)' (first digit run is taken).
    """
    if isinstance(value, (list, tuple)):
        out: List[int] = []
        for v in value:
            out.extend(_parseParkingAlgorithms(v))
        return out
    if isinstance(value, bool):
        raise ValueError(f"Cannot interpret parkingAlgorithm: {value!r}")
    if isinstance(value, int):
        return [value]
    if isinstance(value, float) and value.is_integer():
        return [int(value)]
    if isinstance(value, str):
        match = re.search(r"\d+", value)
        if match:
            return [int(match.group())]
    raise ValueError(f"Cannot interpret parkingAlgorithm: {value!r}")
