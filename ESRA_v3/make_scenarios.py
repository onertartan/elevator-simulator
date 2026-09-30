"""
make_scenarios.py
=================
Generate two additional dispatch snapshots for the CIIS 2026 study, matched
to the snapshot the 157,500-run sweep optimised.

The paper's claim is a FRAMEWORK for a fixed, hard-wired building
(40 floors, 8 cars). The replication that claim needs is therefore not
"another building" but "another traffic realisation of the same building at
the same load", so that the reader can see whether the factor ranking is a
property of the building-and-load or an accident of one particular snapshot.

Everything that defines the search problem is held fixed:

  * 40 floors, 8 cars, park algorithm 1, car initials flag unchanged
  * 17 waiting passengers, split 8 going up / 9 going down
  * 7 distinct up hall calls + 9 distinct down hall calls
    -> nVar = 16 and HC_numofups = 7, so the GA searches a search space of
       exactly the same size and shape (8^16). Changing this would change
       the problem, not the realisation, and the comparison would be void.
  * 9 in-car destinations, one car carrying two of them
  * 4 cars travelling up, 4 travelling down
  * travel-distance distribution: bootstrapped from the original snapshot
    and accepted only when the mean lands within +/-5 % of the original's
    7.35 floors, so difficulty is matched by construction rather than hoped for

Everything that constitutes the realisation is redrawn:

  * car floors, car directions, in-car destination floors
  * passenger origin floors and destination floors

Direction consistency is enforced, not assumed: an upward car's destinations
lie above it, a downward car's below it, an up passenger's destination above
their origin, a down passenger's below.

Output files keep the original workbook's layout cell for cell -- the file is
copied and only the data cells are overwritten -- because the simulator reads
by cell position.
"""
from __future__ import annotations

import argparse
import os
import shutil
from collections import Counter

import numpy as np
import openpyxl

SRC = "initials_files/initials_file.xlsx"
SHEET = "Sayfa1"
FIRST = 3                       # first data row

COL = dict(floors="A", inc="B", inter="C", out="D", cars="E", park="F",
           init="G", car_floor="H", car_state="I", car_df="J",
           up_floor="K", up_df="L", up_car="M",
           dn_floor="N", dn_df="O", dn_car="P")


def read_snapshot(path):
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]

    def col(letter):
        return [ws[f"{letter}{r}"].value
                for r in range(FIRST, ws.max_row + 1)]

    def clean(letter):
        return [v for v in col(letter) if v is not None]

    n_cars = int(ws[f"{COL['cars']}{FIRST}"].value)
    # A blank J cell means this car has no committed destinations; do not
    # remove it and shift all following cars' destination sets left.
    car_df = []
    for row in range(FIRST, FIRST + n_cars):
        value = ws[f"{COL['car_df']}{row}"].value
        car_df.append([] if value is None else [int(x) for x in str(value).split()])

    def passenger_pairs(floor_column, destination_column):
        pairs = []
        for floor, destination in zip(col(floor_column), col(destination_column)):
            if floor is None and destination is None:
                continue
            if floor is None or destination is None:
                raise ValueError("passenger row must contain both origin and destination")
            pairs.append((int(floor), int(destination)))
        return pairs
    s = dict(
        n_floors=ws[f"{COL['floors']}{FIRST}"].value,
        n_cars=ws[f"{COL['cars']}{FIRST}"].value,
        traffic=(ws[f"{COL['inc']}{FIRST}"].value,
                 ws[f"{COL['inter']}{FIRST}"].value,
                 ws[f"{COL['out']}{FIRST}"].value),
        car_floor=[int(v) for v in clean(COL["car_floor"])],
        car_state=[int(v) for v in clean(COL["car_state"])],
        car_df=car_df,
        up=passenger_pairs(COL["up_floor"], COL["up_df"]),
        dn=passenger_pairs(COL["dn_floor"], COL["dn_df"]),
    )
    wb.close()
    return s


def descriptors(s):
    up, dn = s["up"], s["dn"]
    dist = [b - a for a, b in up] + [a - b for a, b in dn]
    n = len(up) + len(dn)
    hc_up = sorted(set(f for f, _ in up))
    hc_dn = sorted(set(f for f, _ in dn))
    car_gap = [abs(cf - min([f for f, _ in up + dn], key=lambda x: abs(x - cf)))
               for cf in s["car_floor"]]
    return {
        "floors": s["n_floors"], "cars": s["n_cars"],
        "passengers": n, "up": len(up), "down": len(dn),
        "hall calls (nVar)": len(hc_up) + len(hc_dn),
        "up hall calls": len(hc_up), "down hall calls": len(hc_dn),
        "in-car destinations": sum(len(d) for d in s["car_df"]),
        "cars going up": sum(1 for v in s["car_state"] if v == 1),
        "mean travel distance": float(np.mean(dist)),
        "median travel distance": float(np.median(dist)),
        "max travel distance": int(np.max(dist)),
        "total travel distance": int(np.sum(dist)),
        "passenger floor span": f"{min(f for f, _ in up + dn)}-"
                                f"{max(f for f, _ in up + dn)}",
        "car floor span": f"{min(s['car_floor'])}-{max(s['car_floor'])}",
        "mean car-to-nearest-call gap": float(np.mean(car_gap)),
        "incoming (from fl. 1)": sum(1 for f, _ in up if f == 1),
        "outgoing (to fl. 1)": sum(1 for _, d in dn if d == 1),
    }


def draw(rng, ref, n_floors, tol=0.05, max_tries=200000, should_cancel=None):
    """Resample one realisation with the reference's structure."""
    up_ref, dn_ref = ref["up"], ref["dn"]
    n_up, n_dn = len(up_ref), len(dn_ref)
    n_up_calls = len(set(f for f, _ in up_ref))
    n_dn_calls = len(set(f for f, _ in dn_ref))
    dist_pool = np.array([b - a for a, b in up_ref]
                         + [a - b for a, b in dn_ref])
    target = dist_pool.mean()
    car_gap_pool = np.array([abs(x - f) for f, ds in
                             zip(ref["car_floor"], ref["car_df"]) for x in ds])
    n_cars = ref["n_cars"]
    n_up_cars = sum(1 for v in ref["car_state"] if v == 1)
    n_incar = sum(len(d) for d in ref["car_df"])
    multi = [i for i, d in enumerate(ref["car_df"]) if len(d) > 1]
    n_multi = len(multi)

    for attempt in range(max_tries):
        if should_cancel and attempt % 128 == 0:
            should_cancel()  # generation caller raises its cancellation exception
        # ---- passengers ------------------------------------------------
        d_up = rng.choice(dist_pool, size=n_up, replace=True)
        d_dn = rng.choice(dist_pool, size=n_dn, replace=True)
        if abs(np.concatenate([d_up, d_dn]).mean() - target) > tol * target:
            continue
        up_floors = rng.choice(np.arange(1, n_floors), size=n_up_calls,
                               replace=False)
        # one call carries two passengers, so that the number of distinct
        # hall calls -- and hence nVar -- matches the reference exactly
        idx = list(range(n_up_calls)) + list(
            rng.choice(n_up_calls, size=n_up - n_up_calls, replace=False))
        up = []
        ok = True
        for k, di in zip(idx, d_up):
            f = int(up_floors[k])
            if f + int(di) > n_floors:
                ok = False
                break
            up.append((f, f + int(di)))
        if not ok or len(set(up)) != n_up:
            continue

        dn_floors = rng.choice(np.arange(2, n_floors + 1), size=n_dn_calls,
                               replace=False)
        dn = []
        for f, di in zip(dn_floors, d_dn):
            f = int(f)
            if f - int(di) < 1:
                ok = False
                break
            dn.append((f, f - int(di)))
        if not ok or len(set(dn)) != n_dn:
            continue
        if len(set(f for f, _ in up)) != n_up_calls:
            continue
        if len(set(f for f, _ in dn)) != n_dn_calls:
            continue

        # ---- cars ------------------------------------------------------
        car_floor = rng.choice(np.arange(1, n_floors + 1), size=n_cars,
                               replace=False)
        # a car on the bottom floor cannot be travelling down, and one on
        # the top floor cannot be travelling up
        forced_up = [i for i, f in enumerate(car_floor) if f == 1]
        forced_dn = [i for i, f in enumerate(car_floor) if f == n_floors]
        if len(forced_up) > n_up_cars or len(forced_dn) > n_cars - n_up_cars:
            continue
        free = [i for i in range(n_cars)
                if i not in forced_up and i not in forced_dn]
        need_up = n_up_cars - len(forced_up)
        pick = rng.permutation(free)
        state = np.full(n_cars, -1)
        for i in forced_up:
            state[i] = 1
        for i in pick[:need_up]:
            state[i] = 1

        # in-car destinations: same total count and same number of cars
        # carrying two, always an upward car as in the reference
        counts = np.ones(n_cars, dtype=int)
        up_cars = [i for i in range(n_cars) if state[i] == 1]
        if len(up_cars) < n_multi:
            continue
        for i in rng.choice(up_cars, size=n_multi, replace=False):
            counts[i] += 1
        if counts.sum() != n_incar:
            continue
        car_df, ok = [], True
        for i in range(n_cars):
            f, st = int(car_floor[i]), int(state[i])
            hi = n_floors - f if st == 1 else f - 1
            pool = car_gap_pool[car_gap_pool <= hi]
            if len(pool) < counts[i]:
                ok = False
                break
            gaps = rng.choice(pool, size=counts[i], replace=False)
            ds = sorted({f + st * int(g) for g in gaps})
            if len(ds) != counts[i] or not all(1 <= x <= n_floors for x in ds):
                ok = False
                break
            car_df.append(ds if st == 1 else sorted(ds, reverse=True))
        if not ok:
            continue

        return dict(n_floors=n_floors, n_cars=n_cars, traffic=ref["traffic"],
                    car_floor=[int(v) for v in car_floor],
                    car_state=[int(v) for v in state],
                    car_df=car_df, up=up, dn=dn)
    raise RuntimeError("could not draw a realisation under the constraints")


def validate(s, ref):
    """Every invariant the simulator or the comparison relies on."""
    errs = []
    if s["n_floors"] != ref["n_floors"] or s["n_cars"] != ref["n_cars"]:
        errs.append("building configuration changed")
    if len(s["up"]) != len(ref["up"]) or len(s["dn"]) != len(ref["dn"]):
        errs.append("passenger counts changed")
    if any(b <= a for a, b in s["up"]):
        errs.append("an 'up' passenger does not travel up")
    if any(b >= a for a, b in s["dn"]):
        errs.append("a 'down' passenger does not travel down")
    for f, st, ds in zip(s["car_floor"], s["car_state"], s["car_df"]):
        if st == 1 and not all(x > f for x in ds):
            errs.append(f"upward car at {f} has destination {ds}")
        if st != 1 and not all(x < f for x in ds):
            errs.append(f"downward car at {f} has destination {ds}")
    if len(set(s["car_floor"])) != len(s["car_floor"]):
        errs.append("two cars on the same floor")
    if any(not (1 <= x <= s["n_floors"])
           for x in s["car_floor"] + [f for f, _ in s["up"] + s["dn"]]
           + [d for _, d in s["up"] + s["dn"]]
           + [x for ds in s["car_df"] for x in ds]):
        errs.append("a floor index is outside 1..n_floors")
    a, b = descriptors(s), descriptors(ref)
    for k in ("hall calls (nVar)", "up hall calls", "down hall calls",
              "in-car destinations", "cars going up"):
        if a[k] != b[k]:
            errs.append(f"{k}: {a[k]} != reference {b[k]}")
    return errs


def write_snapshot(src, dst, s):
    shutil.copyfile(src, dst)              # keep layout and formatting intact
    wb = openpyxl.load_workbook(dst)
    ws = wb[SHEET]
    for letter, values in (
            (COL["car_floor"], s["car_floor"]),
            (COL["car_state"], s["car_state"]),
            (COL["car_df"], [" ".join(str(x) for x in d) if len(d) > 1
                             else d[0] if d else None for d in s["car_df"]]),
            (COL["up_floor"], [f for f, _ in s["up"]]),
            (COL["up_df"], [d for _, d in s["up"]]),
            (COL["dn_floor"], [f for f, _ in s["dn"]]),
            (COL["dn_df"], [d for _, d in s["dn"]])):
        for r in range(FIRST, ws.max_row + 1):
            ws[f"{letter}{r}"] = None
        for i, v in enumerate(values):
            ws[f"{letter}{FIRST + i}"] = v
    wb.save(dst)
    wb.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "veri", SRC))
    ap.add_argument("--out", default=os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "veri"))
    ap.add_argument("--seeds", type=int, nargs="+", default=[20260809, 20260810])
    args = ap.parse_args()

    ref = read_snapshot(args.src)
    rows = {"S1 (original, used by the sweep)": descriptors(ref)}
    print("reference snapshot:", args.src)

    for k, seed in enumerate(args.seeds, start=2):
        rng = np.random.default_rng(seed)
        s = draw(rng, ref, ref["n_floors"])
        errs = validate(s, ref)
        if errs:
            raise SystemExit(f"scenario S{k} failed validation: {errs}")
        dst = os.path.join(args.out, f"initials_file_S{k}.xlsx")
        write_snapshot(args.src, dst, s)
        back = read_snapshot(dst)
        if descriptors(back) != descriptors(s):
            raise SystemExit(f"S{k} did not round-trip through the workbook")
        rows[f"S{k} (seed {seed})"] = descriptors(back)
        print(f"  wrote {dst}  (validated, round-tripped)")

    import pandas as pd
    tab = pd.DataFrame(rows)
    tab.index.name = "descriptor"
    out_csv = os.path.join(args.out, "scenario_descriptors.csv")
    tab.to_csv(out_csv)
    print(f"  wrote {out_csv}\n")
    print(tab.to_string())


if __name__ == "__main__":
    main()
